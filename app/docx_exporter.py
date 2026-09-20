"""
Curriculum Docx Exporter

Generates a beautifully formatted, publication-ready Word Document (.docx)
from a GeneratedCurriculum JSON object, complete with tables, rubrics,
assessments, Bloom learning objectives, and multiple curated resources with pedagogy notes.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Dict, Any, Optional

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def _set_cell_text(cell, text: str, bold: bool = False, size_pt: float = 10, color_rgb: tuple[int, int, int] | None = None):
    """Helper to cleanly format cell text."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)


def generate_curriculum_docx(data: Dict[str, Any], save_to_disk: bool = True) -> bytes:
    """
    Generate a formatted Word Document (.docx) from curriculum data.
    Returns the document bytes and optionally saves to resources/.
    """
    doc = docx.Document()

    # Page setup - 0.75 inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    title = data.get("training_title", "Generated Curriculum")
    training_id = data.get("training_id", "N/A")
    generated_at = data.get("generated_at", "")[:10]

    # ── Document Title ────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"🎓 {title}")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(30, 41, 59)
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_after = Pt(4)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Curriculum Blueprint & Instructional Design Specification | Training ID: {training_id}")
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)
    sub_p.paragraph_format.space_after = Pt(14)

    # ── Overview Summary Table ────────────────────────────────────────────────
    aud_sum = data.get("audience_profile_summary") or {}
    total_hours = sum(m.get("duration", 0) for m in data.get("modules", []))

    summary_table = doc.add_table(rows=3, cols=4)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = [
        ("Learner Level", aud_sum.get("learner_level", "Intermediate")),
        ("Total Duration", f"{total_hours:.1f} Hours"),
        ("Delivery Mode", aud_sum.get("delivery_mode", "BLENDED")),
        ("Language", aud_sum.get("language", "English")),
        ("Education Level", aud_sum.get("education_level", "Bachelor's Degree")),
        ("Work Experience", aud_sum.get("work_experience", "Full-Time Professional")),
    ]

    row_idx = 0
    col_idx = 0
    for label, val in headers:
        cell_lbl = summary_table.cell(row_idx, col_idx)
        cell_val = summary_table.cell(row_idx, col_idx + 1)
        _set_cell_text(cell_lbl, label, bold=True, size_pt=9.5, color_rgb=(71, 85, 105))
        _set_cell_text(cell_val, val, bold=False, size_pt=9.5)
        col_idx += 2
        if col_idx >= 4:
            col_idx = 0
            row_idx += 1

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ── Training Objectives ───────────────────────────────────────────────────
    tp = data.get("training_profile") or {}
    if tp.get("general_objectives"):
        h2 = doc.add_heading("1. Strategic Training Objectives", level=1)
        h2.paragraph_format.space_before = Pt(12)
        for go in tp["general_objectives"]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(go)
            r.font.size = Pt(10.5)

    # ── Audience Profile & Prerequisites ─────────────────────────────────────
    ap = data.get("audience_profile") or {}
    if ap.get("specific_prerequisites") or ap.get("specific_courses"):
        h2 = doc.add_heading("2. Target Audience & Prerequisites", level=1)
        h2.paragraph_format.space_before = Pt(12)
        if ap.get("specific_prerequisites"):
            p_lbl = doc.add_paragraph()
            p_lbl.add_run("Entry Prerequisites:").bold = True
            for pr in ap["specific_prerequisites"]:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(pr).font.size = Pt(10)
        if ap.get("specific_courses"):
            p_lbl = doc.add_paragraph()
            p_lbl.add_run("Recommended Prior Coursework:").bold = True
            for c in ap["specific_courses"]:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(c).font.size = Pt(10)

    # ── Modules & Lessons ────────────────────────────────────────────────────
    h1 = doc.add_heading("3. Detailed Curriculum Structure & Learning Units", level=1)
    h1.paragraph_format.space_before = Pt(14)

    for mod in data.get("modules", []):
        m_order = mod.get("module_order", 1)
        m_name = mod.get("name", "Module")
        m_dur = mod.get("duration", 0)
        m_dur_type = mod.get("duration_type", "HOURS")

        h2 = doc.add_heading(f"Module {m_order}: {m_name} ({m_dur} {m_dur_type})", level=2)
        h2.paragraph_format.space_before = Pt(10)

        p_desc = doc.add_paragraph()
        p_desc.add_run(f"Description: ").bold = True
        p_desc.add_run(mod.get("description", ""))

        p_strat = doc.add_paragraph()
        p_strat.add_run(f"Teaching Strategy: ").bold = True
        p_strat.add_run(f"{mod.get('teaching_strategy', 'N/A')} | ")
        p_strat.add_run(f"Differentiation: ").bold = True
        p_strat.add_run(f"{mod.get('differentiation_strategies', 'N/A')}")

        # Lessons within module
        for l in mod.get("lessons", []):
            l_name = l.get("name", "Lesson")
            l_dur = l.get("duration", 1)
            l_bloom = l.get("bloom_level", "Apply")

            h3 = doc.add_heading(f"• Lesson: {l_name} [{l_bloom} · {l_dur} hrs]", level=3)
            h3.paragraph_format.space_before = Pt(6)

            p_obj = doc.add_paragraph()
            p_obj.add_run("  Learning Objective: ").bold = True
            p_obj.add_run(l.get("objective", ""))

            p_ldesc = doc.add_paragraph()
            p_ldesc.add_run("  Context & Overview: ").bold = True
            p_ldesc.add_run(l.get("description", ""))

            # Curated Multiple Candidate Resources Table with Pedagogy Notes
            media_list = l.get("media_resources", [])
            if media_list:
                p_res_title = doc.add_paragraph()
                r_title = p_res_title.add_run("  Curated Learning Resources & Decision Guide:")
                r_title.bold = True
                r_title.font.size = Pt(9.5)
                p_res_title.paragraph_format.space_after = Pt(2)

                res_table = doc.add_table(rows=len(media_list) + 1, cols=4)
                res_table.alignment = WD_TABLE_ALIGNMENT.CENTER

                # Headers
                col_headers = ["Resource Name & Type", "Role / Pedagogy Notes", "Difficulty & Time", "Direct Access Link"]
                for c_idx, h_text in enumerate(col_headers):
                    _set_cell_text(res_table.cell(0, c_idx), h_text, bold=True, size_pt=8.5, color_rgb=(15, 23, 42))

                for r_idx, res in enumerate(media_list, 1):
                    rec_prefix = "⭐ [TOP PICK] " if res.get("is_primary") else ""
                    res_name = f"{rec_prefix}{res.get('name', 'Resource')} ({res.get('file_type', 'LINK')})"
                    notes = res.get("pedagogy_notes") or res.get("description") or "Curated domain material."
                    timing = f"{res.get('difficulty_level', 'Intermediate')} · {res.get('estimated_time', '20m')}"
                    url = res.get("url") or "Internal Repository"

                    _set_cell_text(res_table.cell(r_idx, 0), res_name, bold=bool(res.get("is_primary")), size_pt=8.5)
                    _set_cell_text(res_table.cell(r_idx, 1), notes, size_pt=8)
                    _set_cell_text(res_table.cell(r_idx, 2), timing, size_pt=8)
                    _set_cell_text(res_table.cell(r_idx, 3), url, size_pt=7.5, color_rgb=(37, 99, 235))

        # Module Rubric
        for rub in mod.get("rubrics", []):
            h3_rub = doc.add_heading(f"Module Rubric: {rub.get('title', 'Rubric')}", level=3)
            h3_rub.paragraph_format.space_before = Pt(8)

            crit_list = rub.get("criteria", [])
            if crit_list:
                rub_table = doc.add_table(rows=len(crit_list) + 1, cols=5)
                rub_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                rub_headers = ["Criterion (Weight)", "Level 4 (Exceeds)", "Level 3 (Meets)", "Level 2 (Developing)", "Level 1 (Beginning)"]
                for c_idx, h_text in enumerate(rub_headers):
                    _set_cell_text(rub_table.cell(0, c_idx), h_text, bold=True, size_pt=8.5, color_rgb=(15, 23, 42))

                for c_idx, crit in enumerate(crit_list, 1):
                    w_pct = round(crit.get("weight", 0) * 100)
                    crit_name = f"{crit.get('criterion','')} ({w_pct}%)\n{crit.get('description','')}"
                    lvls = crit.get("levels", {})
                    _set_cell_text(rub_table.cell(c_idx, 0), crit_name, bold=True, size_pt=8)
                    _set_cell_text(rub_table.cell(c_idx, 1), lvls.get("4", ""), size_pt=8)
                    _set_cell_text(rub_table.cell(c_idx, 2), lvls.get("3", ""), size_pt=8)
                    _set_cell_text(rub_table.cell(c_idx, 3), lvls.get("2", ""), size_pt=8)
                    _set_cell_text(rub_table.cell(c_idx, 4), lvls.get("1", ""), size_pt=8)

        # Knowledge Check Assessment
        for asmt in mod.get("assessments", []):
            h3_asmt = doc.add_heading(f"Assessment: {asmt.get('title', 'Quiz')} ({asmt.get('type','quiz')})", level=3)
            h3_asmt.paragraph_format.space_before = Pt(8)
            p_info = doc.add_paragraph()
            p_info.add_run(f"Duration: {asmt.get('duration_minutes', 20)} mins | Passing Score: {asmt.get('passing_score', 70)}% | Max Attempts: {asmt.get('max_attempts', 1)}").italic = True

            for q_idx, q in enumerate(asmt.get("questions", []), 1):
                p_q = doc.add_paragraph()
                p_q.add_run(f"Q{q_idx} [{q.get('points', 10)} pts]: ").bold = True
                p_q.add_run(q.get("question") or q.get("prompt") or "")
                for opt in q.get("options", []):
                    p_opt = doc.add_paragraph(style="List Bullet")
                    p_opt.add_run(opt).font.size = Pt(9.5)
                if q.get("correct_answer"):
                    p_ans = doc.add_paragraph()
                    p_ans.add_run(f"  Answer Key: {q['correct_answer']}").italic = True

    # ── Export Bytes ──────────────────────────────────────────────────────────
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    # Save to workspace resources/ directory if requested
    if save_to_disk:
        try:
            clean_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
            out_file = Path("resources") / f"{clean_title}.docx"
            out_file.parent.mkdir(parents=True, exist_ok=True)
            out_file.write_bytes(docx_bytes)
        except Exception:
            pass

    return docx_bytes
